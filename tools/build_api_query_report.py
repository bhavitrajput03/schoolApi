from pathlib import Path
import re
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "ET-Teacher-API-SQL-Query-Report.docx"
SOURCE = ROOT / "docs" / "api-sql-query-reference.md"

NAVY = "0B2545"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
CODE_BG = "F6F8FA"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(dxa))
    tcW.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MID_GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, size=9, color=NAVY, bold=True)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trPr.append(repeat)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_font(r, size=8.5, color="1F2937")
    return table


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trPr.append(repeat)
    shade(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + ": ")
    set_font(r, size=10.5, color=NAVY, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.5, color=NAVY)


def code_block(doc, text):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.0
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CODE_BG)
    pPr.append(shd)
    r = p.add_run(text.rstrip())
    set_font(r, name="Consolas", size=8.2, color="172B4D")


def add_inline_markdown(paragraph, text):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Consolas", size=9.5, color="9B1C1C")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=11, color="1F2937", bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=11, color="1F2937")


def markdown_to_doc(doc, markdown):
    lines = markdown.splitlines()
    i = 0
    in_code = False
    code = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                code_block(doc, "\n".join(code))
                code = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
            i += 1
            continue
        if line.startswith("| ") and i + 1 < len(lines) and re.match(r"^\|[- |]+\|$", lines[i + 1]):
            headers = [x.strip() for x in line.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([x.strip().replace("`", "") for x in lines[i].strip("|").split("|")])
                i += 1
            n = len(headers)
            widths = [9360 // n] * n
            widths[-1] += 9360 - sum(widths)
            add_table(doc, headers, rows, widths)
            continue
        if line.startswith("1. ") or re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline_markdown(p, re.sub(r"^\d+\. ", "", line))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, line[2:])
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline_markdown(p, line)
        i += 1
    if code:
        code_block(doc, "\n".join(code))


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before, after, color in (
        ("Heading 1", 16, 18, 10, BLUE),
        ("Heading 2", 13, 14, 7, BLUE),
        ("Heading 3", 12, 10, 5, NAVY),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Code Block" not in [s.name for s in styles]:
        styles.add_style("Code Block", 1)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("ET Teacher API | Technical Reference"), size=9, color=MID_GRAY)
    add_page_number(section.footer.paragraphs[0])

    # Editorial cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("TECHNICAL REPORT"), size=11, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run("ET Teacher API"), size=30, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    set_font(p.add_run("API-to-Query and Database Table Mapping"), size=16, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Complete endpoint inventory, SQL execution flow, table ownership, validation, authorization and write behavior"), size=11, color=MID_GRAY, italic=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(110)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"Generated from the current codebase | {date.today().strftime('%d %B %Y')}"), size=10, color=MID_GRAY)
    doc.add_page_break()

    doc.add_heading("1. Report scope", level=1)
    doc.add_paragraph("This report documents the implemented PHP API exactly as it exists in the current project. It covers the central MySQL school registry, tenant selection, bearer-token authentication, all public and protected routes, SQL Server queries, joins, read/write tables, validation, transactions, upsert keys and response behavior.")
    add_callout(doc, "Important", "MySQL selects which SQL Server database belongs to a school code. All school business data remains in the selected SQL Server database; MySQL does not store students, marks, attendance or API users.")

    doc.add_heading("2. Runtime architecture", level=1)
    steps = [
        "Login receives schoolCode, username and password.",
        "MySQL school_connections resolves schoolCode to connection_string, database_name and encrypted database_password.",
        "CredentialCipher decrypts the SQL Server password in memory.",
        "PDO SQLSRV connects to the selected tenant database.",
        "Login validates ApiUser and creates an ApiAuthToken row.",
        "The access token carries a non-secret school routing prefix; every protected request reselects the correct SQL Server database before running its business query.",
    ]
    for item in steps:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)

    doc.add_heading("2.1 MySQL tenant lookup", level=2)
    code_block(doc, "SELECT connection_string, database_name, database_password\nFROM school_connections\nWHERE school_code = ?\nLIMIT 1;")
    add_table(doc, ["MySQL column", "Purpose"], [
        ("id", "Surrogate primary key."),
        ("school_code", "Unique client/school lookup value supplied at login."),
        ("connection_string", "SQL Server endpoint and User Id; e.g. local SSH tunnel or VPS-local address."),
        ("database_name", "Tenant SQL Server database, such as SchoolManagement."),
        ("database_password", "AES-256-GCM encrypted SQL login password."),
    ], [2200, 7160])

    doc.add_heading("3. Endpoint inventory", level=1)
    endpoints = [
        ("POST", "/api/auth/login", "No", "ApiLoginAttempt, ApiUser, ApiSchool, ApiAuthToken", "Authenticate + issue token"),
        ("POST", "/api/auth/logout", "Yes", "ApiAuthToken", "Revoke current token"),
        ("GET", "/api/auth/me", "Yes", "ApiAuthToken, ApiUser", "Current user"),
        ("POST", "/api/auth/change-password", "Yes", "ApiUser, ApiAuthToken", "Change hash + revoke other sessions"),
        ("GET", "/api/teacher/dashboard", "Yes", "OwnerSession, ApiTeacherAssignment", "Teacher summary"),
        ("GET", "/api/teacher/classes", "Yes", "ApiTeacherAssignment, ClassMaster", "Assigned classes"),
        ("GET", "/api/teacher/classes/{classId}/sections", "Yes", "ApiTeacherAssignment, SectionMaster", "Assigned sections"),
        ("GET", "/api/teacher/classes/{classId}/sections/{sectionId}/subjects", "Yes", "ApiTeacherAssignment, CBSEExamSubject", "Assigned subjects"),
        ("GET", "/api/teacher/classes/{classId}/sections/{sectionId}/students", "Yes", "StudentSession, Student", "Class roster"),
        ("GET", "/api/exams/terms", "Yes", "CBSEExamTerm, CBSEExamTermOption", "Term options"),
        ("GET", "/api/marks/max-marks", "Yes", "Assignments, subjects, marks, roster", "Configured max marks"),
        ("POST", "/api/marks/max-marks", "Yes", "Marks, roster, term/sub-subject", "Bulk max-marks upsert"),
        ("GET", "/api/marks/subject-wise/students", "Yes", "StudentSession, Student, marks", "Roster + marks"),
        ("POST", "/api/marks/subject-wise", "Yes", "StudentSession, marks", "Subject-wise marks upsert"),
        ("GET", "/api/marks/student-wise/{studentId}", "Yes", "Assignments, subjects, marks, roster", "One student across subjects"),
        ("POST", "/api/marks/student-wise", "Yes", "StudentSession, marks", "Student-wise marks upsert"),
        ("GET", "/api/attendance/students", "Yes", "StudentSession, Student, AttItem", "Daily attendance roster"),
        ("POST", "/api/attendance", "Yes", "StudentSession, AttItem", "Attendance upsert"),
    ]
    add_table(doc, ["Method", "Endpoint", "Auth", "Primary tables", "Action"], endpoints, [720, 3000, 650, 2750, 2240])

    doc.add_heading("4. Shared protected-endpoint queries", level=1)
    doc.add_paragraph("Every protected route runs the token lookup before its controller query. Most teacher and marks routes also determine the current academic session and verify the teacher assignment.")
    code_block(doc, "SELECT u.ApiUserID,u.EmployeeID,u.SchoolBranchID,u.DisplayName,u.Username,u.Role\nFROM ApiAuthToken t\nJOIN ApiUser u ON u.ApiUserID=t.ApiUserID\nWHERE t.TokenHash=? AND t.RevokedAt IS NULL\n  AND t.ExpiresAt>SYSUTCDATETIME() AND u.IsActive=1;\n\nUPDATE ApiAuthToken SET LastUsedAt=SYSUTCDATETIME() WHERE TokenHash=?;")
    code_block(doc, "SELECT TOP 1 OwnerSessionID\nFROM OwnerSession\nORDER BY CASE WHEN CAST(GETDATE() AS date) BETWEEN CAST(StartDate AS date)\n  AND CAST(EndDate AS date) THEN 0 ELSE 1 END, EndDate DESC;")
    code_block(doc, "SELECT COUNT(*)\nFROM ApiTeacherAssignment\nWHERE ApiUserID=? AND ClassID=? AND SectionID=?\n  AND OwnerSessionID=? AND IsActive=1\n  AND SubjectID=?; -- subject filter only for subject-specific routes")

    doc.add_page_break()
    doc.add_heading("5. Detailed API and SQL reference", level=1)
    markdown = SOURCE.read_text(encoding="utf-8")
    # Remove duplicated front matter already covered in the report.
    start = markdown.find("## Authentication APIs")
    if start >= 0:
        markdown = markdown[start:]
    markdown_to_doc(doc, markdown)

    doc.add_heading("6. Implementation corrections and exact current behavior", level=1)
    doc.add_heading("6.1 Max marks source", level=2)
    doc.add_paragraph("The current implementation intentionally uses the existing CBSEExamMarksEntry.MaxMarks column. It does not create a separate max-marks configuration table. A max value is represented across the active students' matching marks rows for the same academic session, class, section, term option and mapped sub-subject.")
    doc.add_heading("6.2 GET marks fallback", level=2)
    doc.add_paragraph("Both marks GET flows use COALESCE(current student's MaxMarks, latest configured MaxMarks in the same class/section/session/term/sub-subject). This means a student without a direct marks row can still receive the class-level max value when another roster row contains that configuration.")
    code_block(doc, "CAST(COALESCE(marks.MaxMarks, config.MaxMarks) AS float) AS maxMarks")
    doc.add_heading("6.3 Duplicate prevention", level=2)
    doc.add_paragraph("Marks writes use StudentID + OwnerSessionID + TermOptionID + SubSubjectID as the logical uniqueness key. The transaction takes UPDLOCK and HOLDLOCK before choosing UPDATE versus INSERT. The deployment also includes a unique-index migration for this key.")
    doc.add_heading("6.4 Attendance update key", level=2)
    doc.add_paragraph("Attendance is updated by StudentID plus calendar date. If no row is affected, a new AttItem row is inserted. The current code does not include school session, class or section in the AttItem update predicate; membership is validated separately through StudentSession before the write.")

    doc.add_heading("7. Consolidated table dictionary", level=1)
    rows = [
        ("MySQL.school_connections", "Tenant routing", "R", "schoolCode -> SQL connection/database/password"),
        ("ApiSchool", "API school-code mapping", "R", "Login school code and branch mapping"),
        ("ApiUser", "API identities", "R/W", "Login, profile, password, last login"),
        ("ApiAuthToken", "Sessions", "R/W", "Token validation, issue, last use, revoke"),
        ("ApiLoginAttempt", "Login throttling", "R/W", "15-minute failed-attempt window"),
        ("ApiTeacherAssignment", "Authorization", "R", "Teacher-to-session/class/section/subject scope"),
        ("OwnerSession", "Academic year/session", "R", "Select current session"),
        ("ClassMaster", "Class lookup", "R", "Class name and ordering"),
        ("SectionMaster", "Section lookup", "R", "Section name"),
        ("Student", "Student master", "R", "Identity, admission number, father name"),
        ("StudentSession", "Roster", "R", "Student class/section/session/roll membership"),
        ("CBSEExamTerm", "Exam term", "R", "Term master"),
        ("CBSEExamTermOption", "Term option", "R", "optionId used as termOptionId"),
        ("CBSEExamSubject", "Subject master", "R", "Teacher-facing subject IDs/names"),
        ("CBSEExamSubSubject", "Marks mapping", "R", "First sub-subject selected for each subject"),
        ("CBSEExamMarksEntry", "Marks + max marks", "R/W", "MaxMarks, ObtainMarks, percentage, flags"),
        ("AttItem", "Attendance", "R/W", "Date, status and numeric attendance value"),
    ]
    add_table(doc, ["Table", "Role", "Access", "Used for"], rows, [2600, 2200, 900, 3660])

    doc.add_heading("8. Data rules for the Flutter/app team", level=1)
    rules = [
        "Use optionId from GET /api/exams/terms as termOptionId.",
        "Use subjectId returned by the assigned-subject endpoint; the backend maps it to CBSEExamSubSubjectID.",
        "Null obtainedMarks means no mark entered (or absent/medical), not automatically zero.",
        "A GET may return a roster student with null marks fields because marks are joined with LEFT JOIN.",
        "The server may supply maxMarks from another configured row in the same context via the fallback query.",
        "Never create a second marks entry for the same uniqueness key; POST endpoints already update the existing row.",
        "Attendance statuses are P, A, H/2 and Holiday.",
        "All protected endpoints require Authorization: Bearer <accessToken>.",
    ]
    for rule in rules:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(rule)

    doc.add_heading("9. Operational and security observations", level=1)
    observations = [
        "SQL Server passwords are encrypted in MySQL and decrypted only in application memory.",
        "Prepared statements are used for runtime values; PDO emulated prepares are disabled.",
        "Login responses deliberately do not distinguish an unknown school, username or password.",
        "Password changes use a transaction and revoke all other active tokens.",
        "Marks and attendance batch writes use transactions and rollback on error.",
        "Production must keep APP_DEBUG=false so raw SQL/driver errors are not returned to clients.",
        "The local-to-live development setup requires the SSH tunnel to remain active when connection_string uses 127.0.0.1,14330.",
    ]
    for item in observations:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

    # Core properties
    doc.core_properties.title = "ET Teacher API - API-to-Query and Database Table Mapping"
    doc.core_properties.subject = "Technical API and SQL query reference"
    doc.core_properties.author = "ET Teacher Backend"
    doc.core_properties.keywords = "ET Teacher, API, SQL Server, MySQL, queries, tables"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
